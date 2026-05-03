# from django.shortcuts import render, redirect
# from .models import Client
# from .utils import start_gem_scraping, check_states_in_pdf
# import json
# import os
# from django.http import HttpResponse, JsonResponse
# from openpyxl import Workbook
# from openpyxl.styles import Font
# from django.views.decorators.http import require_http_methods
# from django.views.decorators.http import require_http_methods
# import subprocess
# import sys
# from concurrent.futures import ThreadPoolExecutor, as_completed

# def index(request):
#     # 1. Database se clients fetch karna
#     clients = Client.objects.all().order_by('-created_at')
    
#     # 2. Path aur JSON data loading
#     base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
#     cat_path = os.path.join(base_dir, 'gem_categories.json')
#     state_path = os.path.join(base_dir, 'states.json')

#     # Categories Load Karo
#     try:
#         with open(cat_path, 'r', encoding='utf-8') as f:
#             categories_list = json.load(f)
#     except:
#         categories_list = []

#     # States Load Karo
#     try:
#         with open(state_path, 'r', encoding='utf-8') as f:
#             states_data = json.load(f)
#             states_list = states_data.get('states', [])
#     except:
#         states_list = ["Uttar Pradesh", "Delhi", "Karnataka", "Maharashtra", "Haryana"]

#     # Session se results uthana
#     final_results = request.session.get('results', [])

#     if request.method == "POST":
#         if 'add_client' in request.POST:
#             cats = ", ".join(request.POST.getlist('categories'))
#             states = ", ".join(request.POST.getlist('states'))
#             Client.objects.create(
#                 company_name=request.POST.get('company_name'),
#                 group_name=request.POST.get('group_name'),
#                 categories=cats,
#                 states=states
#             )
#             return redirect('index')

#         elif 'scrape_now' in request.POST:
#             client_id = request.POST.get('client_id')
#             selected_client = Client.objects.get(id=client_id)
#             request.session['selected_client_id'] = client_id
#             cat_list = [c.strip() for c in selected_client.categories.split(',') if c.strip()]
            
#             # Use a single browser session to scrape multiple categories sequentially (lower startup overhead)
#             try:
#                 data = start_gem_scraping_multiple(cat_list, "")
#                 temp_results = data or []
#             except Exception:
#                 # fallback to previous per-category sequential scraping
#                 temp_results = []
#                 for cat in cat_list:
#                     data = start_gem_scraping(cat, "")
#                     if data:
#                         for item in data:
#                             if not any(res.get('bid_no') == item.get('bid_no') for res in temp_results):
#                                 temp_results.append(item)
            
#             request.session['all_results'] = temp_results
#             request.session['results'] = temp_results
#             final_results = temp_results

#         elif 'filter_by_state' in request.POST:
#             # Use 'all_results' backup to ensure we filter from the complete list of tenders
#             results_to_filter = request.session.get('all_results', request.session.get('results', []))
#             client_id = request.POST.get('client_id')
#             selected_client = Client.objects.get(id=client_id)
#             request.session['selected_client_id'] = client_id
            
#             # Client ke assigned states
#             target_states = [s.strip().upper() for s in selected_client.states.split(',') if s.strip()]
            
#             filtered_results = []
#             remaining = []
#             # Fast department check first
#             for bid in results_to_filter:
#                 if any(state in (bid.get('department') or '').upper() for state in target_states):
#                     filtered_results.append(bid)
#                 else:
#                     remaining.append(bid)

#             # Parallelize PDF checks for remaining bids
#             if remaining:
#                 max_workers = min(5, len(remaining))
#                 with ThreadPoolExecutor(max_workers=max_workers) as ex:
#                     future_map = {ex.submit(check_states_in_pdf, bid.get('link', ''), target_states): bid for bid in remaining}
#                     for fut in as_completed(future_map):
#                         bid = future_map[fut]
#                         try:
#                             if fut.result():
#                                 filtered_results.append(bid)
#                         except:
#                             continue
            
#             request.session['results'] = filtered_results
#             final_results = filtered_results

#         # --- LOGIC: Excel Export with Auto-Column Width ---
#         elif 'export_excel' in request.POST:
#             results_to_export = request.session.get('results', [])
#             if results_to_export:
#                 wb = Workbook()
#                 ws = wb.active
#                 ws.title = "GeM Scraping Results"
                
#                 headers = ['Bid Number', 'Items / Category', 'Department & Address', 'Start Date', 'End Date', 'Link']
#                 for col_num, header in enumerate(headers, 1):
#                     cell = ws.cell(row=1, column=col_num, value=header)
#                     cell.font = Font(bold=True)
                
#                 # Writing Data
#                 for row_num, bid in enumerate(results_to_export, 2):
#                     ws.cell(row=row_num, column=1, value=bid.get('bid_no', 'N/A'))
#                     ws.cell(row=row_num, column=2, value=bid.get('items', 'N/A'))
#                     ws.cell(row=row_num, column=3, value=bid.get('department', 'N/A'))
#                     ws.cell(row=row_num, column=4, value=bid.get('start_date', 'N/A'))
#                     ws.cell(row=row_num, column=5, value=bid.get('end_date', 'N/A'))
#                     ws.cell(row=row_num, column=6, value=bid.get('link', 'N/A'))

#                 # --- AUTO-ADJUST COLUMN WIDTH LOGIC ---
#                 for column in ws.columns:
#                     max_length = 0
#                     column_letter = column[0].column_letter # Column name (A, B, C...)
#                     for cell in column:
#                         try:
#                             if len(str(cell.value)) > max_length:
#                                 max_length = len(str(cell.value))
#                         except:
#                             pass
#                     # Thoda buffer (padding) add karke width set karna
#                     adjusted_width = (max_length + 2)
#                     ws.column_dimensions[column_letter].width = adjusted_width

#                 response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
#                 response['Content-Disposition'] = 'attachment; filename="gem_bids_report.xlsx"'
#                 wb.save(response)
#                 return response

#     return render(request, 'index.html', {
#         'clients': clients, 
#         'categories': categories_list, 
#         'states': states_list,
#         'results': final_results
#     })


# # API Endpoint: Get all clients
# @require_http_methods(["GET"])
# def get_all_clients(request):
#     clients = Client.objects.all().order_by('-created_at')
#     clients_data = [
#         {
#             'id': client.id,
#             'company_name': client.company_name,
#             'group_name': client.group_name,
#             'categories': client.categories,
#             'states': client.states,
#             'created_at': client.created_at.strftime("%d-%m-%Y %H:%M")
#         }
#         for client in clients
#     ]
#     return JsonResponse({'clients': clients_data})


# # API Endpoint: Update client
# @require_http_methods(["POST"])
# def update_client(request):
#     try:
#         data = json.loads(request.body)
#         client_id = data.get('id')
#         client = Client.objects.get(id=client_id)
        
#         client.company_name = data.get('company_name', client.company_name)
#         client.group_name = data.get('group_name', client.group_name)
#         client.categories = data.get('categories', client.categories)
#         client.states = data.get('states', client.states)
#         client.save()
        
#         return JsonResponse({'success': True, 'message': 'Client updated successfully'})
#     except Client.DoesNotExist:
#         return JsonResponse({'success': False, 'message': 'Client not found'}, status=404)
#     except Exception as e:
#         return JsonResponse({'success': False, 'message': str(e)}, status=400)


# # API Endpoint: Delete client
# @require_http_methods(["POST"])
# def delete_client(request):
#     try:
#         data = json.loads(request.body)
#         client_id = data.get('id')
#         client = Client.objects.get(id=client_id)
#         client.delete()
        
#         return JsonResponse({'success': True, 'message': 'Client deleted successfully'})
#     except Client.DoesNotExist:
#         return JsonResponse({'success': False, 'message': 'Client not found'}, status=404)
#     except Exception as e:
#         return JsonResponse({'success': False, 'message': str(e)}, status=400)


# # API Endpoint: Save custom category
# @require_http_methods(["POST"])
# def save_custom_category(request):
#     try:
#         data = json.loads(request.body)
#         category = data.get('category', '').strip()
        
#         if not category:
#             return JsonResponse({'success': False, 'message': 'Category cannot be empty'}, status=400)
        
#         base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
#         cat_path = os.path.join(base_dir, 'gem_categories.json')
        
#         # Load existing categories
#         try:
#             with open(cat_path, 'r', encoding='utf-8') as f:
#                 categories_list = json.load(f)
#         except:
#             categories_list = []
        
#         # Add new category if not already exists
#         if category not in categories_list:
#             categories_list.append(category)
            
#             # Save back to JSON file
#             with open(cat_path, 'w', encoding='utf-8') as f:
#                 json.dump(categories_list, f, ensure_ascii=False, indent=4)
            
#             return JsonResponse({'success': True, 'message': f'Category "{category}" saved successfully'})
#         else:
#             return JsonResponse({'success': False, 'message': f'Category "{category}" already exists'}, status=400)
    
#     except Exception as e:
#         return JsonResponse({'success': False, 'message': str(e)}, status=400)


# # API Endpoint: Save custom state
# @require_http_methods(["POST"])
# def save_custom_state(request):
#     try:
#         data = json.loads(request.body)
#         state = data.get('state', '').strip()
        
#         if not state:
#             return JsonResponse({'success': False, 'message': 'State cannot be empty'}, status=400)
        
#         base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
#         state_path = os.path.join(base_dir, 'states.json')
        
#         # Load existing states
#         try:
#             with open(state_path, 'r', encoding='utf-8') as f:
#                 states_data = json.load(f)
#                 states_list = states_data.get('states', [])
#         except:
#             states_list = []
#             states_data = {}
        
#         # Add new state if not already exists
#         if state not in states_list:
#             states_list.append(state)
#             states_data['states'] = states_list
            
#             # Save back to JSON file
#             with open(state_path, 'w', encoding='utf-8') as f:
#                 json.dump(states_data, f, ensure_ascii=False, indent=4)
            
#             return JsonResponse({'success': True, 'message': f'State "{state}" saved successfully'})
#         else:
#             return JsonResponse({'success': False, 'message': f'State "{state}" already exists'}, status=400)
    
#     except Exception as e:
#         return JsonResponse({'success': False, 'message': str(e)}, status=400)


# # API Endpoint: Start background category scraper
# @require_http_methods(["POST"])
# def refresh_categories(request):
#     try:
#         base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
#         script_path = os.path.join(base_dir, 'extractor', 'category.py')

#         # Start the scraper as a background process using the same Python executable
#         # Output is appended to a log file so the server doesn't block.
#         log_file = os.path.join(base_dir, 'category_scrape.log')
#         with open(log_file, 'a', encoding='utf-8') as lf:
#             subprocess.Popen([sys.executable, script_path], cwd=base_dir, stdout=lf, stderr=lf)

#         return JsonResponse({'started': True})
#     except Exception as e:
#         return JsonResponse({'started': False, 'message': str(e)}, status=500)


# # API Endpoint: Get current categories from gem_categories.json
# @require_http_methods(["GET"])
# def get_categories(request):
#     try:
#         base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
#         cat_path = os.path.join(base_dir, 'gem_categories.json')
#         try:
#             with open(cat_path, 'r', encoding='utf-8') as f:
#                 categories_list = json.load(f)
#         except:
#             categories_list = []
#         return JsonResponse({'categories': categories_list})
#     except Exception as e:
#         return JsonResponse({'categories': [], 'message': str(e)}, status=500)






from django.shortcuts import render, redirect
from .models import Client
from .utils import start_gem_scraping, start_gem_scraping_multiple, check_states_in_pdf

import json
import os
from django.http import HttpResponse, JsonResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from django.views.decorators.http import require_http_methods
from concurrent.futures import ThreadPoolExecutor, as_completed
import subprocess
import sys


def index(request):
    clients = Client.objects.all().order_by('-created_at')

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    cat_path = os.path.join(base_dir, 'gem_categories.json')
    state_path = os.path.join(base_dir, 'states.json')

    # Load categories
    try:
        with open(cat_path, 'r', encoding='utf-8') as f:
            categories_list = json.load(f)
    except:
        categories_list = []

    # Load states
    try:
        with open(state_path, 'r', encoding='utf-8') as f:
            states_data = json.load(f)
            states_list = states_data.get('states', [])
    except:
        states_list = ["Uttar Pradesh", "Delhi", "Haryana"]

    final_results = request.session.get('results', [])

    if request.method == "POST":

        # ---------------- ADD CLIENT ---------------- #
        if 'add_client' in request.POST:
            cats = ", ".join(request.POST.getlist('categories'))
            states = ", ".join(request.POST.getlist('states'))

            Client.objects.create(
                company_name=request.POST.get('company_name'),
                group_name=request.POST.get('group_name'),
                categories=cats,
                states=states
            )
            return redirect('index')

        # ---------------- SCRAPE ---------------- #
        elif 'scrape_now' in request.POST:
            client_id = request.POST.get('client_id')
            selected_client = Client.objects.get(id=client_id)

            cat_list = [c.strip() for c in selected_client.categories.split(',') if c.strip()]

            try:
                temp_results = start_gem_scraping_multiple(cat_list, "")
            except Exception as e:
                print("Multi scrape error:", e)

                temp_results = []
                for cat in cat_list:
                    data = start_gem_scraping(cat, "")
                    if data:
                        for item in data:
                            if not any(res.get('bid_no') == item.get('bid_no') for res in temp_results):
                                temp_results.append(item)

            request.session['all_results'] = temp_results
            request.session['results'] = temp_results
            final_results = temp_results

        # ---------------- FILTER BY STATE ---------------- #
        elif 'filter_by_state' in request.POST:
            results_to_filter = request.session.get('all_results', [])

            client_id = request.POST.get('client_id')
            selected_client = Client.objects.get(id=client_id)

            target_states = [s.strip().upper() for s in selected_client.states.split(',') if s.strip()]

            filtered_results = []
            remaining = []

            # Fast filter (department)
            for bid in results_to_filter:
                if any(state in (bid.get('department') or '').upper() for state in target_states):
                    filtered_results.append(bid)
                else:
                    remaining.append(bid)

            # PDF filter (parallel)
            if remaining:
                with ThreadPoolExecutor(max_workers=5) as executor:
                    future_map = {
                        executor.submit(check_states_in_pdf, bid.get('link', ''), target_states): bid
                        for bid in remaining
                    }

                    for future in as_completed(future_map):
                        bid = future_map[future]
                        try:
                            if future.result():
                                filtered_results.append(bid)
                        except:
                            continue

            request.session['results'] = filtered_results
            final_results = filtered_results

        # ---------------- EXPORT EXCEL ---------------- #
        elif 'export_excel' in request.POST:
            results_to_export = request.session.get('results', [])

            if results_to_export:
                wb = Workbook()
                ws = wb.active
                ws.title = "GeM Results"

                headers = ['Bid No', 'Items', 'Department', 'Start Date', 'End Date', 'Link']

                for col, header in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=col, value=header)
                    cell.font = Font(bold=True)

                for row_num, bid in enumerate(results_to_export, 2):
                    ws.cell(row=row_num, column=1, value=bid.get('bid_no'))
                    ws.cell(row=row_num, column=2, value=bid.get('items'))
                    ws.cell(row=row_num, column=3, value=bid.get('department'))
                    ws.cell(row=row_num, column=4, value=bid.get('start_date'))
                    ws.cell(row=row_num, column=5, value=bid.get('end_date'))
                    ws.cell(row=row_num, column=6, value=bid.get('link'))

                # Auto width
                for col in ws.columns:
                    max_length = 0
                    col_letter = col[0].column_letter

                    for cell in col:
                        try:
                            if cell.value:
                                max_length = max(max_length, len(str(cell.value)))
                        except:
                            pass

                    ws.column_dimensions[col_letter].width = max_length + 2

                response = HttpResponse(
                    content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
                )
                response['Content-Disposition'] = 'attachment; filename="gem_results.xlsx"'

                wb.save(response)
                return response

    return render(request, 'index.html', {
        'clients': clients,
        'categories': categories_list,
        'states': states_list,
        'results': final_results
    })


# ---------------- API ---------------- #

@require_http_methods(["GET"])
def get_all_clients(request):
    clients = Client.objects.all().order_by('-created_at')

    data = [{
        "id": c.id,
        "company_name": c.company_name,
        "group_name": c.group_name,
        "categories": c.categories,
        "states": c.states
    } for c in clients]

    return JsonResponse({"clients": data})


@require_http_methods(["POST"])
def delete_client(request):
    data = json.loads(request.body)
    client = Client.objects.get(id=data.get("id"))
    client.delete()

    return JsonResponse({"success": True})


@require_http_methods(["POST"])
def update_client(request):
    data = json.loads(request.body)
    client = Client.objects.get(id=data.get("id"))

    client.company_name = data.get("company_name")
    client.group_name = data.get("group_name")
    client.categories = data.get("categories")
    client.states = data.get("states")
    client.save()

    return JsonResponse({"success": True})


@require_http_methods(["POST"])
def refresh_categories(request):
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    script_path = os.path.join(base_dir, 'extractor', 'category.py')

    subprocess.Popen([sys.executable, script_path], cwd=base_dir)

    return JsonResponse({"started": True})


@require_http_methods(["GET"])
def get_categories(request):
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    cat_path = os.path.join(base_dir, 'gem_categories.json')

    try:
        with open(cat_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except:
        data = []

    return JsonResponse({"categories": data})